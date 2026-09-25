#!/usr/bin/env python3
"""Build the tax-credit CSVs (EITC and Child Tax Credit) for all four levels.

Sources, in data/raw/tax_credits/ (see the README there for URLs):

  * 23incyallnoagi_hi.csv -- the Hawaii rows (state total and four counties)
    of IRS Statistics of Income (SOI) county data for tax year 2023, used as
    published for the state and the counties.
  * 22zp12hi.xlsx -- SOI ZIP code data for Hawaii, tax year 2022, the latest
    SOI has published. Its ZIP rows are allocated to legislative districts
    (below); 22incyallnoagi_hi.csv, the 2022 county rows, checks that.
  * act107_earnedincome_txcredit_2023_tables.xlsx -- the Hawaii Department of
    Taxation's report on the state EITC for tax year 2023, which checks the
    state EITC estimate (below).

plus data/raw/crosswalks/geocorr2022_zcta_to_{sldl22,sldu22,county}.csv --
Geocorr 2022 (Missouri Census Data Center): the share of each ZIP code
tabulation area's 2020 population in each 2022 House / Senate district (and
county).

Measures (SOI field names):

    total_returns   N1       number of returns
    eitc_returns    N59660   returns with the earned income credit
    federal EITC    A59660   earned income credit amount
    ctc_returns     N11070   returns with the refundable child tax credit
                             (additional child tax credit)
    CTC amount      A11070   refundable child tax credit amount

The dashboard's CTC is the refundable part, the one that reaches families
with little or no income tax. (The earlier files took the CTC amount from the
next column over, "earned income credit with one qualifying child", which
halved every average.)

State EITC: from tax year 2023 Hawaii's credit is 40% of the federal EITC,
refundable (Acts 114, SLH 2022 and 163, SLH 2023; before, 20% and
nonrefundable), so it is estimated as 40% of the federal credit, on 2023
returns, the first year at that rate. Statewide the estimate comes within 2%
of the average new claim in the Department of Taxation's report for the same
year ($950 against $938); prorated credits for part-year residents pull the
actual a little lower. The script checks this, so a change in the rate (Act
163's 40% is set to lapse after 2027) shows up when YEAR moves on. The
estimated total (40% of all federal EITC dollars) runs about 10% above the
credits actually claimed, since not every federal claimant claims the state
credit; the dashboard shows only the average.

Districts: each ZIP's 2022 totals are split across districts by the Geocorr
factors, so districts add up to the state. SOI folds ZIP codes with few
returns, and nonresidential ones such as Honolulu PO boxes, into "other"
(99999); that row is spread over the ZIP code areas that have no row of their
own, by population. Its rates are close to the state's, and two rural House
districts (5 and 17) have about half their people in those areas, so their
figures lean toward the statewide mix. The ZIP-based county totals come
within 0.25 points of SOI's county data for the same year on every rate;
the script fails past 0.5. The districts in each county (House and Senate
districts each lie within one) are then scaled to that county's 2023 totals,
measure by measure, so they add up to the counties, as
scripts/build_alice.py does. Once SOI publishes ZIP code data for YEAR, set
ZIP_YEAR to match.

Outputs data/processed/tax_credits/hawaii_{state,county,house_district,
senate_district}_tax_credits_<YEAR>.csv, with rates as fractions (the data
loader converts them to percentages). Then run
scripts/build_static/02_build_layer_geojsons.py to rebuild the map layers.
"""
from __future__ import annotations

import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / 'data' / 'raw' / 'tax_credits'
CROSSWALKS = ROOT / 'data' / 'raw' / 'crosswalks'
OUT = ROOT / 'data' / 'processed' / 'tax_credits'
YEAR = 2023      # SOI county data and the state EITC report
ZIP_YEAR = 2022  # SOI ZIP code data, for the district pattern within counties
STATE_EITC_SHARE = 0.40

MEASURES = ('returns', 'eitc_n', 'eitc_a', 'ctc_n', 'ctc_a')
COUNTY_NAMES = {'15001': 'Hawaii', '15003': 'Honolulu', '15007': 'Kauai', '15009': 'Maui'}
KALAWAO, MAUI = '15005', '15009'
COLUMNS = [
    'NAME', 'state', 'district', 'geoid', 'total_returns', 'eitc_returns', 'eitc_participation_rate',
    'federal_eitc_avg_amount', 'federal_eitc_total_amount', 'state_eitc_avg_amount',
    'state_eitc_total_amount', 'ctc_returns', 'ctc_participation_rate', 'ctc_avg_amount',
    'ctc_total_amount', 'total_population',
]


def read_zip_data() -> tuple[dict, dict]:
    """State total and per-ZIP totals from the SOI ZIP workbook (amounts in
    thousands of dollars, as published)."""
    rows = list(openpyxl.load_workbook(RAW / f'{ZIP_YEAR % 100}zp12hi.xlsx', read_only=True,
                                       data_only=True).active.iter_rows(values_only=True))
    header = [re.sub(r'\s+', ' ', str(h or '')).strip() for h in rows[3]]
    units = [str(h or '').strip() for h in rows[4]]

    def count_col(title: str) -> int:
        i = next(i for i, h in enumerate(header) if h.startswith(title))
        assert units[i] == 'Number of returns' and units[i + 1] == 'Amount', (title, units[i:i + 2])
        return i

    cols = {'returns': 2}
    cols['eitc_n'] = count_col('Earned income credit [')
    cols['ctc_n'] = count_col('Refundable child tax credit or additional child tax credit')
    cols['eitc_a'], cols['ctc_a'] = cols['eitc_n'] + 1, cols['ctc_n'] + 1

    state, zips = None, {}
    for r in rows[6:]:
        if not isinstance(r[0], (int, float)):
            continue
        values = {m: float(r[i] or 0) for m, i in cols.items()}
        if r[0] == 0 and r[1] == 'Total':
            state = values
        elif r[0] != 0 and r[1] is None:  # a ZIP's all-incomes row
            zips[f'{int(r[0]):05d}'] = values
    assert state and zips, 'SOI ZIP workbook layout changed'
    return state, zips


def read_county_data(year: int) -> tuple[dict, dict]:
    """State total and per-county totals from the Hawaii rows of SOI's county
    data (amounts in thousands of dollars)."""
    fields = {'returns': 'N1', 'eitc_n': 'N59660', 'eitc_a': 'A59660', 'ctc_n': 'N11070', 'ctc_a': 'A11070'}
    with open(RAW / f'{year % 100}incyallnoagi_hi.csv', newline='') as f:
        rows = {f"15{int(r['COUNTYFIPS']):03d}": {m: float(r[k]) for m, k in fields.items()}
                for r in csv.DictReader(f)}
    state = rows.pop('15000')
    assert set(rows) == set(COUNTY_NAMES), f'SOI county rows changed: {sorted(rows)}'
    return state, rows


def read_state_eitc_claims() -> tuple[float, float]:
    """Number and dollar amount of new state EITC claims for YEAR, from Table 1
    of the Department of Taxation's EITC report (the rest of that table is
    earlier years' nonrefundable credits carried forward)."""
    wb = openpyxl.load_workbook(RAW / f'act107_earnedincome_txcredit_{YEAR}_tables.xlsx', read_only=True,
                                data_only=True)
    rows = list(wb['Table 1'].iter_rows(values_only=True))
    assert rows[2][0] == f'Tax Year {YEAR}', rows[2]
    header = next(r for r in rows if r[0] == 'Federal Adjusted Gross Income Range')
    i = next(i for i, h in enumerate(header) if h and h.startswith('Total New Credit Claimed'))
    total = next(r for r in rows if r[0] == 'Total')
    return float(total[i]), float(total[i + 1])


def read_crosswalk(target: str) -> list[tuple[str, str, float, float]]:
    """(zcta, target, afact, pop20) rows; zcta is '' for blocks outside any ZCTA."""
    with open(CROSSWALKS / f'geocorr2022_zcta_to_{target}.csv', newline='') as f:
        rows = list(csv.reader(f))[2:]  # a names row, then a labels row
    return [(z.strip(), geo.strip(), float(row[-1]), float(row[-2])) for z, geo, *row in rows]


def allocate(zips: dict, crosswalk: list) -> tuple[dict, dict]:
    """Totals and 2020 population per target area."""
    other = dict(zips['99999'])
    for z, values in zips.items():  # any SOI ZIP without a ZCTA joins "other"
        if z != '99999' and not any(row[0] == z for row in crosswalk):
            for m in MEASURES:
                other[m] += values[m]
    unlisted = [(geo, pop) for z, geo, _, pop in crosswalk if z not in zips]
    unlisted_pop = sum(pop for _, pop in unlisted)

    totals = defaultdict(lambda: dict.fromkeys(MEASURES, 0.0))
    population = defaultdict(float)
    for z, geo, afact, pop in crosswalk:
        population[geo] += pop
        if z in zips:
            for m in MEASURES:
                totals[geo][m] += afact * zips[z][m]
    for geo, pop in unlisted:
        for m in MEASURES:
            totals[geo][m] += pop / unlisted_pop * other[m]
    return totals, population


def district_counties(crosswalk: list, county_crosswalk: list) -> dict:
    """Each district's county, from where its population lives."""
    zcta_county = {z: MAUI if fips == KALAWAO else fips
                   for z, fips, afact, _ in county_crosswalk if z and afact > 0.5}
    pop = defaultdict(lambda: defaultdict(float))
    for z, d, _, p in crosswalk:
        if z:
            pop[d][zcta_county[z]] += p
    out = {}
    for d, by_county in pop.items():
        fips, top = max(by_county.items(), key=lambda kv: kv[1])
        assert top >= 0.99 * sum(by_county.values()), f'district {d} spans counties: {dict(by_county)}'
        out[d] = fips
    return out


def scale_to_counties(totals: dict, county_of: dict, counties: dict) -> dict:
    """The districts in each county, scaled measure by measure to add up to
    that county's totals."""
    sums = defaultdict(lambda: dict.fromkeys(MEASURES, 0.0))
    for d, t in totals.items():
        for m in MEASURES:
            sums[county_of[d]][m] += t[m]
    return {d: {m: t[m] * counties[county_of[d]][m] / sums[county_of[d]][m] for m in MEASURES}
            for d, t in totals.items()}


def num(x: float, digits: int) -> float | int:
    """Rounded, and written without a trailing .0 when whole (SOI counts are)."""
    x = round(x, digits)
    return int(x) if float(x).is_integer() else x


def row_for(t: dict, **ids) -> dict:
    federal = t['eitc_a'] * 1000
    ctc = t['ctc_a'] * 1000
    return {
        **ids,
        'total_returns': num(t['returns'], 1),
        'eitc_returns': num(t['eitc_n'], 1),
        'eitc_participation_rate': num(t['eitc_n'] / t['returns'], 6),
        'federal_eitc_avg_amount': num(federal / t['eitc_n'], 2),
        'federal_eitc_total_amount': num(federal, 0),
        'state_eitc_avg_amount': num(STATE_EITC_SHARE * federal / t['eitc_n'], 2),
        'state_eitc_total_amount': num(STATE_EITC_SHARE * federal, 0),
        'ctc_returns': num(t['ctc_n'], 1),
        'ctc_participation_rate': num(t['ctc_n'] / t['returns'], 6),
        'ctc_avg_amount': num(ctc / t['ctc_n'], 2),
        'ctc_total_amount': num(ctc, 0),
    }


def write(level: str, rows: list[dict]) -> None:
    path = OUT / f'hawaii_{level}_tax_credits_{YEAR}.csv'
    columns = [c for c in COLUMNS if c != 'district' or level.endswith('district')]
    with open(path, 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=columns, lineterminator='\n')
        w.writeheader()
        w.writerows(rows)
    print(f'  wrote {path.relative_to(ROOT) if path.is_relative_to(ROOT) else path} ({len(rows)} rows)')


def rates(t: dict) -> tuple[float, float]:
    return 100 * t['eitc_n'] / t['returns'], 100 * t['ctc_n'] / t['returns']


def main() -> int:
    zip_state, zips = read_zip_data()
    county_crosswalk = read_crosswalk('county')
    county_zip_totals, county_pop = allocate(zips, county_crosswalk)
    county_pop[MAUI] += county_pop.pop(KALAWAO, 0)  # Kalawao is mapped with Maui

    # The ZIP-based county totals should agree with SOI's own county data.
    for fips, t in read_county_data(ZIP_YEAR)[1].items():
        worst = max(abs(a - b) for a, b in zip(rates(t), rates(county_zip_totals[fips])))
        if worst > 0.5:
            print(f'  ! {fips}: ZIP-based rates differ from SOI county data by {worst:.2f} points')
            return 1

    # The state EITC estimate should match the average new claim the
    # Department of Taxation reports.
    state, counties = read_county_data(YEAR)
    claims, amount = read_state_eitc_claims()
    estimate = STATE_EITC_SHARE * 1000 * state['eitc_a'] / state['eitc_n']
    print(f'  state EITC {YEAR}: estimated ${estimate:,.0f} per claim, reported ${amount / claims:,.0f}')
    if abs(estimate / (amount / claims) - 1) > 0.02:
        print('  ! the state EITC estimate is off by more than 2%; check STATE_EITC_SHARE')
        return 1

    OUT.mkdir(parents=True, exist_ok=True)
    write('state', [row_for(state, NAME='Hawaii (Statewide)', state=15, geoid='15',
                            total_population=round(sum(county_pop.values())))])
    write('county', [row_for(counties[fips], NAME=COUNTY_NAMES[fips], state=15, geoid=fips,
                             total_population=round(county_pop[fips])) for fips in sorted(counties)])
    for level, target, chamber in [('house_district', 'sldl22', 'House'), ('senate_district', 'sldu22', 'Senate')]:
        crosswalk = read_crosswalk(target)
        totals, population = allocate(zips, crosswalk)
        for m in MEASURES:  # districts add up to the state (within SOI rounding)
            assert abs(sum(t[m] for t in totals.values()) - zip_state[m]) <= 0.001 * zip_state[m] + 200, (level, m)
        totals = scale_to_counties(totals, district_counties(crosswalk, county_crosswalk), counties)
        write(level, [
            row_for(totals[d], NAME=f'State {chamber} District {int(d):02d}; Hawaii; Hawaii', state=15,
                    district=int(d), geoid=f'15{int(d):02d}', total_population=round(population[d]))
            for d in sorted(totals, key=int)
        ])
    return 0


if __name__ == '__main__':
    sys.exit(main())
