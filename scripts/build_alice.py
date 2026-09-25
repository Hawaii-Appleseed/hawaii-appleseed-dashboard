#!/usr/bin/env python3
"""Build the ALICE workbook the data pipeline reads, for all four levels.

The dashboard's ALICE figure is the share of households that are ALICE:
above the federal poverty level but below the cost of basics (the ALICE
Threshold). Households in poverty are not included; that is the separate
poverty rate. The earlier workbook mixed two measures: its state and county
sheets held this ALICE-only share, but its House and Senate sheets held the
share below the ALICE Threshold (poverty plus ALICE), about 10 points higher,
so districts looked far worse than the counties they make up.

Source: United For ALICE's Hawaii data sheet from its 2026 report, in
data/raw/alice/ (see the README there), for data year 2024:

  * State and counties: the county figures as published; the state is their
    sum. Kalawao (40 households) is counted with Maui, as on the map.
  * House and Senate districts: the ZIP code figures, split across districts
    by the Geocorr 2022 ZIP-to-district factors (2020 population) in
    data/raw/crosswalks/. The ZIP figures are 5-year estimates while the
    county ones are 1-year, and ZIP codes under 100 households aren't
    reported, so the districts in each county are then scaled to that
    county's published totals (households, poverty and ALICE households
    separately), as scripts/build_medicaid_2024.py does. House and Senate
    districts lie within one county each, so every district aggregate
    matches its county, and the counties the state.

Writes data/ALICE By Geography (<year>).xlsx with sheets State, Counties,
House and Senate (see "alice" in src/config/data_sources.json): an area
column, alice_rate (percent), and the household counts behind it. Then run
scripts/build_static/02_build_layer_geojsons.py to rebuild the map layers.
"""
from __future__ import annotations

import csv
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'data' / 'raw' / 'alice' / '2026 ALICE - Hawaii Data Sheet.xlsx'
CROSSWALKS = ROOT / 'data' / 'raw' / 'crosswalks'
YEAR = 2024
OUT = ROOT / 'data' / f'ALICE By Geography ({YEAR}).xlsx'

MEASURES = ('households', 'poverty', 'alice')
COUNTY_NAMES = {'15001': 'Hawaii', '15003': 'Honolulu', '15007': 'Kauai', '15009': 'Maui'}
KALAWAO, MAUI = '15005', '15009'


def sheet_rows(wb, name: str) -> list[dict]:
    rows = list(wb[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:] if any(v is not None for v in r)]


def read_source() -> tuple[dict, dict]:
    """County and ZIP code (ZCTA) household counts for YEAR."""
    wb = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    counties = defaultdict(lambda: dict.fromkeys(MEASURES, 0.0))
    for r in sheet_rows(wb, 'County'):
        if r['Year'] == YEAR:
            fips = MAUI if str(r['GEO.id2']) == KALAWAO else str(r['GEO.id2'])
            for m, col in zip(MEASURES, ('Households', 'Poverty Households', 'ALICE Households')):
                counties[fips][m] += float(r[col])
    zips = {
        str(r['GEO.id2']).split('_')[0]: {
            m: float(r[col]) for m, col in zip(MEASURES, ('Households', 'Poverty.Households', 'ALICE.Households'))
        }
        for r in sheet_rows(wb, 'Subcounty') if r['Type'] == 'Zip_Code' and r['Year'] == YEAR
    }
    assert set(counties) == set(COUNTY_NAMES) and zips, 'ALICE data sheet layout changed'
    return dict(counties), zips


def read_crosswalk(target: str) -> list[tuple[str, str, float, float]]:
    """(zcta, target, afact, pop20) rows; zcta is '' for blocks outside any ZCTA."""
    with open(CROSSWALKS / f'geocorr2022_zcta_to_{target}.csv', newline='') as f:
        rows = list(csv.reader(f))[2:]  # a names row, then a labels row
    return [(z.strip(), geo.strip(), float(row[-1]), float(row[-2])) for z, geo, *row in rows]


def district_counties(crosswalk: list) -> dict:
    """Each district's county, from where its population lives."""
    zcta_county = {}
    for z, fips, afact, _ in read_crosswalk('county'):
        if z and afact > 0.5:
            zcta_county[z] = MAUI if fips == KALAWAO else fips
    pop = defaultdict(lambda: defaultdict(float))
    for z, d, _, p in crosswalk:
        if z:
            pop[d][zcta_county[z]] += p
    out = {}
    for d, by_county in pop.items():
        fips, top = max(by_county.items(), key=lambda kv: kv[1])
        assert top >= 0.99 * sum(by_county.values()), f'district {d} spans counties: {dict(by_county)}'
        out[d] = fips
    return out


def districts(zips: dict, counties: dict, target: str) -> dict:
    crosswalk = read_crosswalk(target)
    raw = defaultdict(lambda: dict.fromkeys(MEASURES, 0.0))
    for z, d, afact, _ in crosswalk:
        if z in zips:
            for m in MEASURES:
                raw[d][m] += afact * zips[z][m]
    county_of = district_counties(crosswalk)
    sums = defaultdict(lambda: dict.fromkeys(MEASURES, 0.0))
    for d, t in raw.items():
        for m in MEASURES:
            sums[county_of[d]][m] += t[m]
    return {d: {m: t[m] * counties[county_of[d]][m] / sums[county_of[d]][m] for m in MEASURES}
            for d, t in raw.items()}


def row(name, t: dict) -> list:
    return [name, round(100 * t['alice'] / t['households'], 1),
            round(t['households']), round(t['poverty']), round(t['alice'])]


def build() -> dict[str, list[list]]:
    """Sheet name -> rows, header first."""
    counties, zips = read_source()
    state = {m: sum(t[m] for t in counties.values()) for m in MEASURES}
    counts = ['alice_rate', 'Households', 'Poverty Households', 'ALICE Households']
    house = districts(zips, counties, 'sldl22')
    senate = districts(zips, counties, 'sldu22')
    return {
        'State': [['State', *counts], row('Hawaii', state)],
        'Counties': [['County', *counts]] + [row(COUNTY_NAMES[f], counties[f]) for f in sorted(counties)],
        'House': [['District', *counts]] + [row(int(d), house[d]) for d in sorted(house, key=int)],
        'Senate': [['Senate District', *counts]] + [row(int(d), senate[d]) for d in sorted(senate, key=int)],
    }


def main() -> int:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in build().items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        print(f'  {name}: {len(rows) - 1} rows, ALICE {min(r[1] for r in rows[1:])}–{max(r[1] for r in rows[1:])}%')
    wb.save(OUT)
    print(f'  wrote {OUT.relative_to(ROOT)}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
